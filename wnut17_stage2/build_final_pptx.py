"""Comprehensive comparison deck (Bahasa Indonesia, black-only, no colored boxes).
Section A (official split) vs Section B (pooled) vs Bi-LSTM CRF + frozen RoBERTa.
Black titles/headers only, grayscale figures. Writes NER_WNUT17_Perbandingan_Kelompok1.pptx.
"""
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from wnut_data import load
from evalx import evaluate

BLACK = RGBColor(0, 0, 0); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x40, 0x40, 0x40); LGRAY = RGBColor(0xF0, 0xF0, 0xF0)
FT_H = "Cambria"; FT_B = "Calibri"

# numbers
testA = load("test")
E = lambda ev: round(ev["entity"]["f1"] * 100, 1); S = lambda ev: round(ev["surface"]["f1"] * 100, 1)
A_crf = json.load(open("results/to_crf.json"))["test"]; A_svm = json.load(open("results/to_svm.json"))["test"]
A_bl = evaluate(testA, json.load(open("results/pred_to_bilstm_base_nobias_test.json")))
A_ens = json.load(open("results/ensemble_6model.json"))["test_delta0"]
RB = json.load(open("results/bilstm_roberta.json")); RB0, RBd = RB["test_delta0"], RB["test"]
B = json.load(open("results_partB/partB_analysis.json")); st = json.load(open("data_partB/split_stats.json"))

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.text_frame.word_wrap = True
    return tb.text_frame

def setp(p, text, size, bold=False, color=BLACK, font=FT_B, align=PP_ALIGN.LEFT):
    p.text = text; p.alignment = align
    r = p.runs[0]; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font

def title(s, t, sub=None):
    tf = box(s, 0.55, 0.35, 12.2, 1.0); setp(tf.paragraphs[0], t, 30, True, BLACK, FT_H)
    # thin black rule
    ln = s.shapes.add_shape(1, Inches(0.6), Inches(1.25), Inches(12.1), Pt(2))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLACK; ln.line.fill.background()
    if sub:
        tf2 = box(s, 0.55, 1.32, 12.2, 0.5); setp(tf2.paragraphs[0], sub, 14, False, GRAY, FT_B)

def bullets(s, items, x, y, w, h, size=16, gap=6):
    tf = box(s, x, y, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(it, tuple):
            lvl, txt = (it[0], it[1]) if len(it) == 2 else (0, it[0])
        else:
            lvl, txt = 0, it
        setp(p, ("- " if lvl == 0 else "    - ") + txt, size - lvl * 2, False, BLACK, FT_B)
        p.space_after = Pt(gap)

def table(s, rows, x, y, w, colw=None, header=True, size=13):
    nr, nc = len(rows), len(rows[0])
    g = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(0.4 * nr)).table
    if colw:
        for j, cw in enumerate(colw): g.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = g.cell(i, j); c.text = str(val); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_top = Pt(2); c.margin_bottom = Pt(2)
            pr = c.text_frame.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
            rr = pr.runs[0]; rr.font.size = Pt(size); rr.font.name = FT_B
            if header and i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = BLACK; rr.font.color.rgb = WHITE; rr.font.bold = True
            else:
                c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else LGRAY; rr.font.color.rgb = BLACK
    return g

def pic(s, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    s.shapes.add_picture(path, Inches(x), Inches(y), **kw)

def foot(s, n):
    tf = box(s, 11.7, 7.05, 1.4, 0.35); setp(tf.paragraphs[0], f"Kelompok 1  |  {n}", 9, False, GRAY, FT_B, PP_ALIGN.RIGHT)

# 1 Title
s = slide()
tf = box(s, 0.8, 2.4, 11.7, 2.2)
setp(tf.paragraphs[0], "NER WNUT17: Perbandingan Tiga Pendekatan", 40, True, BLACK, FT_H)
p = tf.add_paragraph(); setp(p, "Split resmi (Section A)  |  Bi-LSTM CRF + RoBERTa beku  |  Split gabungan 80/10/10 (Section B)", 18, False, GRAY, FT_B); p.space_before = Pt(14)
p = tf.add_paragraph(); setp(p, "Kelompok 1  -  Pemrosesan Bahasa Alami Lanjut  -  Assignment #2 Part 2", 15, False, BLACK, FT_B); p.space_before = Pt(24)

# 2 Ringkasan
s = slide(); title(s, "Ringkasan", "Tiga pendekatan, satu dataset, satu kesimpulan yang jujur")
bullets(s, [
    "Section A (split resmi): tiga model wajib (CRF, SVM, Bi-LSTM CRF) pada split yang sengaja dipisah domain. Ini hasil utama dan jujur.",
    "Peningkatan RoBERTa: menambah fitur RoBERTa beku ke Bi-LSTM CRF pada split yang sama. Naik ~13 poin pada F1 entity dan surface.",
    "Section B (split gabungan 80/10/10): ablasi. Menunjukkan berapa banyak kesulitan berasal dari entitas baru, bukan model yang lebih baik.",
    ("F1 entity test: CRF 16.8, SVM 29.8, Bi-LSTM 40.4, ensemble 41.35  ->  + RoBERTa 54.1  |  Section B: 40.6 / 42.1 / 51.8", ),
], 0.7, 1.7, 12.0, 4.5, size=17, gap=12)
foot(s, 2)

# 3 Dataset & metrik
s = slide(); title(s, "Dataset dan metrik", "WNUT17 emerging entities, 6 tipe, skema IOB2")
bullets(s, [
    "Split resmi dipisah domain dengan sengaja: train = Twitter 2010, dev = YouTube 2017, test = Reddit/StackExchange/Twitter 2017.",
    "Sekitar 98% surface entitas di test tidak pernah muncul di train. Tugasnya mengenali entitas BARU, bukan menghafal.",
    "Dua metrik resmi, keduanya level entity: F1 entity (seqeval, batas dan tipe persis) dan F1 surface (pasangan unik surface-tipe).",
    "Semua model: seed 42, dipilih di dev, test dinilai sekali. Interval bootstrap 2000-resample dilaporkan di test.",
], 0.7, 1.7, 12.0, 4.5, size=17, gap=12)
foot(s, 3)

# 4 Section A hasil
s = slide(); title(s, "Section A: split resmi (hasil utama)", "Dilatih pada file train Twitter saja; fitur final")
table(s, [
    ["Model", "P", "R", "F1 entity", "F1 surface"],
    ["CRF", f"{A_crf['entity']['precision']*100:.1f}", f"{A_crf['entity']['recall']*100:.1f}", f"{E(A_crf)}", f"{S(A_crf)}"],
    ["SVM", f"{A_svm['entity']['precision']*100:.1f}", f"{A_svm['entity']['recall']*100:.1f}", f"{E(A_svm)}", f"{S(A_svm)}"],
    ["Bi-LSTM CRF (baseline)", f"{A_bl['entity']['precision']*100:.1f}", f"{A_bl['entity']['recall']*100:.1f}", f"{E(A_bl)}", f"{S(A_bl)}"],
    ["Bi-LSTM CRF (ensemble 6)", f"{A_ens['entity']['precision']*100:.1f}", f"{A_ens['entity']['recall']*100:.1f}", f"{E(A_ens)}", f"{S(A_ens)}"],
], 0.9, 1.9, 8.0, colw=[3.2, 1.2, 1.2, 1.4, 1.4], size=14)
bullets(s, [
    "CRF dan SVM rendah (16.8 / 29.8): fitur identitas kata jarang menyala saat 98% surface baru.",
    "Ensemble 6 varian seed-42 satu-satunya yang naik di dev DAN test.",
    "Enlarge ke 9 varian (41.1) atau ensemble 5-seed (41.1): F1 surface tetap 39.6.",
    "Kesimpulan: ensembling (reduksi varians) sudah mentok.",
], 0.9, 4.5, 11.5, 2.4, size=15, gap=8)
foot(s, 4)

# 5 RoBERTa metode + hasil
s = slide(); title(s, "Peningkatan: Bi-LSTM CRF + RoBERTa beku", "Split resmi yang sama, seed 42, RoBERTa TIDAK di-fine-tune")
bullets(s, [
    "Tambah satu input beku per token: vektor RoBERTa-base 768d (hidden terakhir, subword di-mean-pool per token, konteks per kalimat).",
    "Arsitektur tetap Bi-LSTM CRF; RoBERTa hanya ekstraktor fitur beku, bukan model ke-4. Delta O-bias dipilih di dev.",
], 0.7, 1.65, 12.0, 1.5, size=16, gap=8)
table(s, [
    ["Bi-LSTM CRF + RoBERTa", "P", "R", "F1 entity", "F1 surface", "CI entity 95%"],
    ["delta 0", f"{RB0['entity']['precision']*100:.1f}", f"{RB0['entity']['recall']*100:.1f}", f"{E(RB0)}", f"{S(RB0)}", "[49.2, 55.7]"],
    ["dev-tuned delta 1.0", f"{RBd['entity']['precision']*100:.1f}", f"{RBd['entity']['recall']*100:.1f}", f"{E(RBd)}", f"{S(RBd)}", "[51.0, 57.2]"],
], 0.9, 3.2, 8.2, colw=[3.0, 1.0, 1.0, 1.4, 1.4, 2.2], size=14)
pic(s, "report_figs_final/f_path.png", 9.0, 2.9, w=4.0)
bullets(s, [
    "Naik ~13 poin pada KEDUA metrik. Batas bawah CI (51.0) jauh di atas 41.35: efek nyata, bukan noise.",
    "Jujur: RoBERTa model kontekstual 2019, tak dimiliki sistem shared task 2017; bukan menang setara. Tetap Twitter-only.",
], 0.9, 5.2, 8.0, 1.8, size=13, gap=6)
foot(s, 5)

# 6 Section B
s = slide(); title(s, "Section B: split gabungan 80/10/10 (ablasi)", "Mixing diizinkan untuk section ini saja; tidak sebanding dengan shared task")
table(s, [
    ["Model", "F1 A (resmi)", "F1 B (gabung)", "surface A", "surface B", "kebocoran test"],
    ["CRF", f"{E(A_crf)}", f"{B['CRF']['entity']['f1']:.1f}", f"{S(A_crf)}", f"{B['CRF']['surface_f1']:.1f}", f"2% -> {st['leakage']['test']['leak_pct']}%"],
    ["SVM", f"{E(A_svm)}", f"{B['SVM']['entity']['f1']:.1f}", f"{S(A_svm)}", f"{B['SVM']['surface_f1']:.1f}", f"2% -> {st['leakage']['test']['leak_pct']}%"],
    ["Bi-LSTM CRF", f"{E(A_bl)}", f"{B['Bi-LSTM CRF']['entity']['f1']:.1f}", f"{S(A_bl)}", f"{B['Bi-LSTM CRF']['surface_f1']:.1f}", f"2% -> {st['leakage']['test']['leak_pct']}%"],
], 0.9, 1.9, 8.2, colw=[2.6, 1.6, 1.6, 1.5, 1.5, 2.2], size=13)
pic(s, "report_figs_partB/f_partB_leaked_recall.png", 9.1, 1.9, w=3.9)
bullets(s, [
    "Konservasi persis (5690/101857/3890), pelanggaran BIO 0, split per dokumen seed 42.",
    "Kebocoran entitas naik ~2% -> 21.1%. Kenaikan F1 terkonsentrasi pada entitas bocor (CRF: 64.8% vs 23.8% pada yang benar-benar baru).",
    "Jadi kenaikan = kebocoran + cakupan fitur, BUKAN model lebih baik. Tidak sebanding dengan leaderboard.",
], 0.9, 4.7, 12.0, 2.2, size=14, gap=7)
foot(s, 6)

# 7 Perbandingan menyeluruh
s = slide(); title(s, "Perbandingan menyeluruh", "Bi-LSTM CRF: split sulit vs split mudah vs RoBERTa")
pic(s, "report_figs_final/f_threeway.png", 0.8, 1.7, w=7.4)
bullets(s, [
    "Bi-LSTM di split resmi (sulit): ~40-41.",
    "Split gabungan (mudah): 51.8, hampir seluruhnya dari kebocoran.",
    "RoBERTa di split SULIT: 54.1, menyamai/melebihi angka split mudah tanpa merusak split jujur.",
    ("Artinya: embedding kontekstual benar-benar MENYELESAIKAN kesulitan WNUT17; split gabungan hanya MENYEMBUNYIKANNYA.", ),
], 8.4, 1.8, 4.6, 4.8, size=15, gap=12)
foot(s, 7)

# 8 Kesimpulan
s = slide(); title(s, "Kesimpulan", None)
bullets(s, [
    "Section A adalah hasil utama dan jujur: CRF 16.8, SVM 29.8, Bi-LSTM 40.4, ensemble 41.35 F1 entity.",
    "Ensembling tidak bisa lebih tinggi karena hanya mengurangi varians (F1 surface mentok 39.6).",
    "RoBERTa beku menaikkan Bi-LSTM CRF ke 54.1 entity / 52.9 surface: +~13 poin nyata, tanpa keluar dari split resmi, tanpa model ke-4.",
    "Section B (ablasi) menjelaskan mengapa angka resmi rendah: mayoritas kesulitan WNUT17 adalah kebaruan entitas; split acak hanya menyembunyikannya di balik kebocoran 21%.",
    ("Rekomendasi: Section A sebagai baseline jujur, RoBERTa sebagai peningkatan, Section B sebagai ablasi (ditandai tidak sebanding dengan shared task).", ),
], 0.7, 1.7, 12.2, 5.0, size=17, gap=14)
foot(s, 8)

prs.save("NER_WNUT17_Perbandingan_Kelompok1.pptx")
print("NER_WNUT17_Perbandingan_Kelompok1.pptx written |", len(prs.slides._sldIdLst), "slides", flush=True)
